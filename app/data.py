"""Data loading and normalization helpers for local processed datasets."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from app.config import DATA_FILES, LEGACY_RAW_DIR, PROCESSED_DIR


def data_path(filename: str) -> Path:
    """Prefer cleaned data, but support the earlier raw-folder layout."""
    for folder in (PROCESSED_DIR, LEGACY_RAW_DIR):
        candidate = folder / filename
        if candidate.exists():
            return candidate
    return PROCESSED_DIR / filename


def load_csv(filename: str, parse_dates: tuple[str, ...] = ()) -> pd.DataFrame:
    path = data_path(filename)
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, parse_dates=list(parse_dates))


def coerce_numeric(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    df = df.copy()
    for col in columns:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def coerce_boolean(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    df = df.copy()
    for col in columns:
        if col in df.columns:
            df[col] = df[col].astype(str).str.lower().map({"true": True, "false": False}).fillna(False)
    return df


def load_data() -> dict[str, pd.DataFrame]:
    data = {
        "master_usd": load_csv(DATA_FILES["master_usd"]),
        "master_idx": load_csv(DATA_FILES["master_idx"]),
        "ffpi_monthly": load_csv(DATA_FILES["ffpi_monthly"], ("date",)),
        "ffpi_annual": load_csv(DATA_FILES["ffpi_annual"]),
        "ghi": load_csv(DATA_FILES["ghi"]),
        "worldbank": load_csv(DATA_FILES["worldbank"]),
    }

    numeric_columns = [
        "year",
        "value",
        "ffpi_food",
        "ffpi_meat",
        "ffpi_dairy",
        "ffpi_cereals",
        "ffpi_oils",
        "ffpi_sugar",
        "food_import_pct",
        "ghi_2000",
        "ghi_2008",
        "ghi_2016",
        "ghi_2025",
    ]
    for key, df in data.items():
        df = coerce_numeric(df, numeric_columns)
        data[key] = coerce_boolean(df, ["is_imputed", "is_outlier"])
    return data


def require_data(data: dict[str, pd.DataFrame]) -> bool:
    missing = [name for name, df in data.items() if df.empty]
    if missing:
        st.error(
            "Missing local data files: "
            + ", ".join(DATA_FILES[name] for name in missing)
            + f". Expected them under {PROCESSED_DIR}."
        )
        return False
    return True


def safe_float(value: object) -> float | None:
    if isinstance(value, str):
        value = value.replace("<", "").strip()
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_ghi_indicators() -> pd.DataFrame:
    path = data_path("global_hunger_index.xlsx")
    if not path.exists():
        return pd.DataFrame()

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["GHI Indicator Values 2025"]
    rows = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        country = row[0]
        if country is None or len(str(country)) > 60:
            continue
        rows.append(
            {
                "country_ghi": country,
                "undernourishment_2024": safe_float(row[4]),
                "child_wasting_2024": safe_float(row[11]),
                "child_stunting_2024": safe_float(row[16]),
                "child_mortality_2023": safe_float(row[20]),
            }
        )
    return pd.DataFrame(rows)
