"""Streamlit dashboard for FAOSTAT food price shock analysis."""

from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook
from plotly.subplots import make_subplots


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

PROCESSED_DIR = ROOT_DIR / "data" / "processed"
LEGACY_RAW_DIR = ROOT_DIR / "data" / "raw"

DATA_FILES = {
    "master_usd": "master_producer_prices_usd.csv",
    "master_idx": "master_producer_price_index.csv",
    "ffpi_monthly": "ffpi_monthly.csv",
    "ffpi_annual": "ffpi_annual.csv",
    "ghi": "ghi_cleaned.csv",
    "worldbank": "worldbank_food_import_pct.csv",
}

CRISIS_PERIODS = [
    (2007, 2009, "2007-09 food crisis"),
    (2010, 2012, "2010-11 Arab Spring"),
    (2021, 2023, "2022 Ukraine war"),
]

COLORS = {
    "AUS": "#14b8a6",
    "NZL": "#3b82f6",
    "FFPI": "#d97706",
    "risk": "#e11d48",
    "crisis": "#f59e0b",
    "neutral": "#64748b",
    "low": "#dbeafe",
    "mid": "#f59e0b",
    "high": "#be123c",
    "highlight": "#fff7ed",
}

CHART_TEMPLATE = "plotly_white"
COUNTRY_COLORS = {
    "AUS": COLORS["AUS"],
    "NZL": COLORS["NZL"],
    "Australia": COLORS["AUS"],
    "New Zealand": COLORS["NZL"],
}
FFPI_COLORS = {
    "ffpi_food": "#d97706",
    "ffpi_cereals": "#e11d48",
    "ffpi_meat": "#14b8a6",
    "ffpi_dairy": "#3b82f6",
    "ffpi_oils": "#8b5cf6",
    "ffpi_sugar": "#f59e0b",
}
RISK_SCALE = [
    [0.00, "#dbeafe"],
    [0.35, "#14b8a6"],
    [0.60, "#f59e0b"],
    [0.82, "#fb7185"],
    [1.00, "#be123c"],
]

EXPOSURE_WEIGHTS = {
    "ghi": 0.6,
    "import_dependency": 0.4,
}
SCENARIO_ASSUMPTIONS = {
    "basket_share": 0.45,
    "transmission_coeff": 0.429,
    "pass_through_rate": 0.50,
}


def apply_chart_style(fig: go.Figure, height: int | None = None, show_legend: bool | None = None) -> go.Figure:
    layout = {
        "template": CHART_TEMPLATE,
        "paper_bgcolor": "rgba(0,0,0,0)",
        "plot_bgcolor": "rgba(0,0,0,0)",
        "margin": {"l": 30, "r": 30, "t": 70, "b": 40},
        "legend_title_text": "",
    }
    if height is not None:
        layout["height"] = height
    if show_legend is not None:
        layout["showlegend"] = show_legend
    fig.update_layout(**layout)
    fig.update_xaxes(gridcolor="rgba(127,127,127,0.22)", zerolinecolor="rgba(127,127,127,0.25)")
    fig.update_yaxes(gridcolor="rgba(127,127,127,0.22)", zerolinecolor="rgba(127,127,127,0.25)")
    return fig


def story_insight(step: str, insight: str, action: str | None = None) -> None:
    text = f"**{step}.** {insight}"
    if action:
        text += f"\n\n**Policy use:** {action}"
    st.info(text)


def data_path(filename: str) -> Path:
    """Prefer cleaned data, but support the earlier raw-folder layout."""
    for folder in (PROCESSED_DIR, LEGACY_RAW_DIR):
        candidate = folder / filename
        if candidate.exists():
            return candidate
    return PROCESSED_DIR / filename


@st.cache_data(show_spinner=False)
def load_csv(filename: str, parse_dates: tuple[str, ...] = ()) -> pd.DataFrame:
    path = data_path(filename)
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, parse_dates=list(parse_dates))


def coerce_numeric(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    df = df.copy()
    for col in columns:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def coerce_boolean(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    df = df.copy()
    for col in columns:
        if col in df.columns:
            df[col] = df[col].astype(str).str.lower().map({"true": True, "false": False}).fillna(False)
    return df


@st.cache_data(show_spinner=False)
def load_data() -> dict[str, pd.DataFrame]:
    data = {
        "master_usd": load_csv(DATA_FILES["master_usd"]),
        "master_idx": load_csv(DATA_FILES["master_idx"]),
        "ffpi_monthly": load_csv(DATA_FILES["ffpi_monthly"], ("date",)),
        "ffpi_annual": load_csv(DATA_FILES["ffpi_annual"]),
        "ghi": load_csv(DATA_FILES["ghi"]),
        "worldbank": load_csv(DATA_FILES["worldbank"]),
    }

    numeric_columns = [
        "year",
        "value",
        "ffpi_food",
        "ffpi_meat",
        "ffpi_dairy",
        "ffpi_cereals",
        "ffpi_oils",
        "ffpi_sugar",
        "food_import_pct",
        "ghi_2000",
        "ghi_2008",
        "ghi_2016",
        "ghi_2025",
    ]
    for key, df in data.items():
        df = coerce_numeric(df, numeric_columns)
        data[key] = coerce_boolean(df, ["is_imputed", "is_outlier"])
    return data


def require_data(data: dict[str, pd.DataFrame]) -> bool:
    missing = [name for name, df in data.items() if df.empty]
    if missing:
        st.error(
            "Missing local data files: "
            + ", ".join(DATA_FILES[name] for name in missing)
            + f". Expected them under {PROCESSED_DIR}."
        )
        return False
    return True


def add_crisis_bands(fig: go.Figure, yref: str = "paper") -> None:
    for start, end, label in CRISIS_PERIODS:
        fig.add_vrect(
            x0=start,
            x1=end,
            fillcolor=COLORS["crisis"],
            opacity=0.08,
            line_width=0,
            annotation_text=label,
            annotation_position="top left",
            annotation_font_size=10,
            yref=yref,
        )


def minmax_scale(series: pd.Series) -> pd.Series:
    min_value = series.min()
    max_value = series.max()
    if pd.isna(min_value) or pd.isna(max_value) or min_value == max_value:
        return pd.Series(0.0, index=series.index)
    return (series - min_value) / (max_value - min_value)


def safe_float(value: object) -> float | None:
    if isinstance(value, str):
        value = value.replace("<", "").strip()
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


@st.cache_data(show_spinner=False)
def load_ghi_indicators() -> pd.DataFrame:
    path = data_path("global_hunger_index.xlsx")
    if not path.exists():
        return pd.DataFrame()

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["GHI Indicator Values 2025"]
    rows = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        country = row[0]
        if country is None or len(str(country)) > 60:
            continue
        rows.append(
            {
                "country_ghi": country,
                "undernourishment_2024": safe_float(row[4]),
                "child_wasting_2024": safe_float(row[11]),
                "child_stunting_2024": safe_float(row[16]),
                "child_mortality_2023": safe_float(row[20]),
            }
        )
    return pd.DataFrame(rows)


def build_vulnerability(
    ghi: pd.DataFrame,
    worldbank: pd.DataFrame,
    hunger_weight: float = EXPOSURE_WEIGHTS["ghi"],
) -> pd.DataFrame:
    import_weight = 1 - hunger_weight
    ghi_slim = ghi[["iso3", "country_ghi", "ghi_2000", "ghi_2008", "ghi_2016", "ghi_2025"]].dropna(
        subset=["iso3", "ghi_2025"]
    )
    latest_imports = (
        worldbank.dropna(subset=["food_import_pct"])
        .sort_values("year", ascending=False)
        .groupby("iso3", as_index=False)
        .first()[["iso3", "country_wb", "year", "food_import_pct"]]
    )
    vuln = ghi_slim.merge(latest_imports, on="iso3", how="inner")
    vuln["ghi_norm"] = minmax_scale(vuln["ghi_2025"])
    vuln["import_dependency_norm"] = minmax_scale(vuln["food_import_pct"])
    vuln["exposure_index"] = (
        hunger_weight * vuln["ghi_norm"] + import_weight * vuln["import_dependency_norm"]
    ) * 100
    return vuln.sort_values("exposure_index", ascending=False)


def build_undernourishment_exposure(
    ghi: pd.DataFrame,
    worldbank: pd.DataFrame,
    hunger_weight: float = EXPOSURE_WEIGHTS["ghi"],
) -> pd.DataFrame:
    import_weight = 1 - hunger_weight
    indicators = load_ghi_indicators()
    if indicators.empty:
        return build_vulnerability(ghi, worldbank, hunger_weight)

    ghi_slim = ghi[["iso3", "country_ghi", "ghi_2000", "ghi_2008", "ghi_2016", "ghi_2025"]].dropna(
        subset=["iso3", "ghi_2025"]
    )
    latest_imports = (
        worldbank.dropna(subset=["food_import_pct"])
        .sort_values("year", ascending=False)
        .groupby("iso3", as_index=False)
        .first()[["iso3", "country_wb", "year", "food_import_pct"]]
    )
    exposure = ghi_slim.merge(indicators, on="country_ghi", how="left").merge(latest_imports, on="iso3", how="inner")
    hunger_source = "undernourishment_2024"
    if exposure[hunger_source].notna().sum() < 10:
        exposure[hunger_source] = exposure["ghi_2025"]
        hunger_source = "ghi_2025"

    exposure[hunger_source] = exposure[hunger_source].fillna(exposure[hunger_source].median())
    exposure["hunger_metric"] = exposure[hunger_source]
    exposure["hunger_metric_label"] = (
        "Undernourishment 2022-24 (% population)" if hunger_source == "undernourishment_2024" else "GHI score 2025"
    )
    exposure["hunger_norm"] = minmax_scale(exposure["hunger_metric"])
    exposure["import_dependency_norm"] = minmax_scale(exposure["food_import_pct"])
    exposure["exposure_index"] = (
        hunger_weight * exposure["hunger_norm"]
        + import_weight * exposure["import_dependency_norm"]
    ) * 100
    return exposure.sort_values("exposure_index", ascending=False)


def zscore(series: pd.Series) -> pd.Series:
    std = series.std()
    if pd.isna(std) or std == 0:
        return pd.Series(0.0, index=series.index)
    return (series - series.mean()) / std


def build_corrected_vulnerability(
    ghi: pd.DataFrame,
    worldbank: pd.DataFrame,
    hunger_weight: float = EXPOSURE_WEIGHTS["ghi"],
) -> pd.DataFrame:
    import_weight = 1 - hunger_weight
    exposure = build_undernourishment_exposure(ghi, worldbank, hunger_weight).copy()
    hunger = exposure["undernourishment_2024"] if "undernourishment_2024" in exposure else exposure["hunger_metric"]
    exposure["undernourishment_2024"] = hunger.fillna(hunger.median())
    exposure["food_import_pct"] = exposure["food_import_pct"].fillna(exposure["food_import_pct"].median())
    exposure["z_undernourishment"] = zscore(exposure["undernourishment_2024"])
    exposure["z_food_import"] = zscore(exposure["food_import_pct"])
    exposure["vulnerability_score"] = (
        hunger_weight * exposure["z_undernourishment"]
        + import_weight * exposure["z_food_import"]
    )
    return exposure.sort_values("vulnerability_score", ascending=False)


def build_ffpi_correlation(master_idx: pd.DataFrame, ffpi_annual: pd.DataFrame) -> pd.DataFrame:
    price_idx = (
        master_idx.groupby(["iso3", "country", "year"], as_index=False)["value"]
        .mean()
        .rename(columns={"value": "avg_producer_price_index"})
    )
    ffpi_cols = ["ffpi_food", "ffpi_cereals", "ffpi_meat", "ffpi_oils", "ffpi_sugar", "ffpi_dairy"]
    merged = price_idx.merge(ffpi_annual[["year", *ffpi_cols]], on="year", how="inner")

    rows = []
    for country, group in merged.groupby("country"):
        for col in ffpi_cols:
            rows.append(
                {
                    "country": country,
                    "ffpi_index": col.replace("ffpi_", "").title(),
                    "correlation": group["avg_producer_price_index"].corr(group[col]),
                }
            )
    return pd.DataFrame(rows).dropna(subset=["correlation"])


def build_wheat_signal(master_usd: pd.DataFrame, ffpi_annual: pd.DataFrame) -> pd.DataFrame:
    wheat = master_usd[master_usd["item"].eq("Wheat")].copy()
    if wheat.empty:
        return pd.DataFrame()
    wheat = wheat.merge(ffpi_annual[["year", "ffpi_food"]], on="year", how="left", suffixes=("", "_annual"))
    if "ffpi_food_annual" in wheat.columns:
        wheat["ffpi_food"] = wheat["ffpi_food_annual"].fillna(wheat["ffpi_food"])
    return wheat.sort_values(["country", "year"])


def build_scenario(
    exposure: pd.DataFrame,
    shock_pct: int | float,
    pass_through_rate: int | float | None = None,
    basket_share: int | float | None = None,
    transmission_coeff: int | float | None = None,
) -> tuple[pd.DataFrame, float]:
    scenario = exposure.copy()
    if pass_through_rate is None:
        pass_through_rate = SCENARIO_ASSUMPTIONS["pass_through_rate"]
    if basket_share is None:
        basket_share = SCENARIO_ASSUMPTIONS["basket_share"]
    if transmission_coeff is None:
        transmission_coeff = SCENARIO_ASSUMPTIONS["transmission_coeff"]
    pass_through_rate = float(pass_through_rate)
    basket_share = float(basket_share)
    transmission_coeff = float(transmission_coeff)
    effective_global_pressure_pct = (
        shock_pct
        * basket_share
        * transmission_coeff
        * pass_through_rate
    )
    scenario["implied_import_cost_pressure_pct"] = (
        scenario["food_import_pct"] * effective_global_pressure_pct / 100
    )
    scenario["scenario_pressure_score"] = scenario["exposure_index"] * scenario["implied_import_cost_pressure_pct"] / 100
    return scenario.sort_values("implied_import_cost_pressure_pct", ascending=False), effective_global_pressure_pct


def assign_policy_action(row: pd.Series) -> str:
    hunger = row.get("ghi_2025", row.get("hunger_metric", 0))
    imports = row.get("food_import_pct", 0)
    if hunger >= 35 or row.get("exposure_index", 0) >= 70:
        return "Preparedness package"
    if imports >= 25:
        return "Import-buffer monitoring"
    return "Watchlist"


def priority_table(exposure: pd.DataFrame, top_n: int = 8) -> pd.DataFrame:
    table = exposure.copy()
    if "exposure_index" not in table and "vulnerability_score" in table:
        table["exposure_index"] = minmax_scale(table["vulnerability_score"]) * 100
    table["policy_action"] = table.apply(assign_policy_action, axis=1)
    score_col = "scenario_pressure_score" if "scenario_pressure_score" in table else "exposure_index"
    return (
        table.nlargest(top_n, score_col)[
            ["country_ghi", "ghi_2025", "food_import_pct", "exposure_index", "policy_action"]
        ]
        .round({"ghi_2025": 1, "food_import_pct": 1, "exposure_index": 1})
        .rename(
            columns={
                "country_ghi": "Country",
                "ghi_2025": "GHI 2025",
                "food_import_pct": "Food imports %",
                "exposure_index": "Exposure index",
                "policy_action": "Suggested policy use",
            }
        )
    )


def build_weight_sensitivity(ghi: pd.DataFrame, worldbank: pd.DataFrame, top_n: int = 10) -> pd.DataFrame:
    rows = []
    for hunger_weight in [0.4, 0.5, 0.6, 0.7, 0.8]:
        exposure = build_corrected_vulnerability(ghi, worldbank, hunger_weight).copy()
        exposure["rank"] = exposure["exposure_index"].rank(ascending=False, method="min")
        for _, row in exposure.nsmallest(top_n, "rank").iterrows():
            rows.append(
                {
                    "country": row["country_ghi"],
                    "food_access_weight": hunger_weight,
                    "import_weight": 1 - hunger_weight,
                    "rank": int(row["rank"]),
                    "exposure_index": row["exposure_index"],
                }
            )
    sensitivity = pd.DataFrame(rows)
    if sensitivity.empty:
        return sensitivity
    summary = (
        sensitivity.groupby("country", as_index=False)
        .agg(
            appearances=("food_access_weight", "count"),
            best_rank=("rank", "min"),
            average_rank=("rank", "mean"),
            average_exposure=("exposure_index", "mean"),
        )
        .sort_values(["appearances", "best_rank", "average_rank"], ascending=[False, True, True])
    )
    return summary


def sidebar_controls(data: dict[str, pd.DataFrame]) -> dict[str, object]:
    ffpi_cols = ["ffpi_food", "ffpi_cereals", "ffpi_meat", "ffpi_dairy", "ffpi_oils", "ffpi_sugar"]
    master = data["master_usd"]
    ffpi = data["ffpi_annual"]

    st.sidebar.title("Food Shock Early Warning")
    st.sidebar.caption("Analysis scope and model assumptions")

    with st.sidebar.expander("Global Analysis Scope", expanded=True):
        year_range = st.slider(
            "Year range",
            min_value=int(ffpi["year"].min()),
            max_value=int(ffpi["year"].max()),
            value=(max(1991, int(ffpi["year"].min())), int(ffpi["year"].max())),
            step=1,
        )
        selected_indices = st.multiselect(
            "FAO Food Price Index series",
            ffpi_cols,
            default=["ffpi_food", "ffpi_cereals", "ffpi_oils"],
            format_func=lambda value: value.replace("ffpi_", "").title(),
        )
        if not selected_indices:
            selected_indices = ["ffpi_food"]

    with st.sidebar.expander("Producer Signal Scope", expanded=True):
        items = sorted(master["item"].dropna().unique())
        default_item = "Wheat" if "Wheat" in items else items[0]
        selected_item = st.selectbox("Commodity", items, index=items.index(default_item))
        selected_countries = st.multiselect(
            "Producer countries",
            ["Australia", "New Zealand"],
            default=["Australia", "New Zealand"],
        )
        if not selected_countries:
            selected_countries = ["Australia", "New Zealand"]

    with st.sidebar.expander("Exposure Assumptions", expanded=True):
        hunger_weight = st.slider(
            "Food-access stress weight",
            min_value=0.0,
            max_value=1.0,
            value=EXPOSURE_WEIGHTS["ghi"],
            step=0.05,
            format="%.2f",
        )
        st.caption(f"Food import dependency weight: {1 - hunger_weight:.2f}")

    with st.sidebar.expander("Scenario Assumptions", expanded=True):
        shock_pct = st.slider("Producer price shock", min_value=0, max_value=60, value=20, step=5, format="+%d%%")
        basket_share_pct = st.slider(
            "Relevant import basket share",
            min_value=0,
            max_value=100,
            value=int(SCENARIO_ASSUMPTIONS["basket_share"] * 100),
            step=5,
            format="%d%%",
        )
        transmission_pct = st.slider(
            "Transmission coefficient",
            min_value=0,
            max_value=100,
            value=int(round(SCENARIO_ASSUMPTIONS["transmission_coeff"] * 100)),
            step=5,
            format="%d%%",
        )
        pass_through_pct = st.slider(
            "Producer-to-import pass-through",
            min_value=0,
            max_value=100,
            value=int(SCENARIO_ASSUMPTIONS["pass_through_rate"] * 100),
            step=5,
            format="%d%%",
        )

    with st.sidebar.expander("Evidence Options", expanded=False):
        dataset_label = st.selectbox(
            "Dataset",
            [
                "Producer prices, USD",
                "Producer price index",
                "FFPI annual",
                "FFPI monthly",
                "Global Hunger Index",
                "World Bank food imports",
            ],
        )
        volatility_country = st.radio("Volatility country", ["Australia", "New Zealand"], horizontal=True)

    return {
        "year_range": year_range,
        "selected_indices": selected_indices,
        "selected_item": selected_item,
        "selected_countries": selected_countries,
        "hunger_weight": hunger_weight,
        "shock_pct": shock_pct,
        "basket_share": basket_share_pct / 100,
        "transmission_coeff": transmission_pct / 100,
        "pass_through_rate": pass_through_pct / 100,
        "dataset_label": dataset_label,
        "volatility_country": volatility_country,
    }


def executive_brief(data: dict[str, pd.DataFrame], controls: dict[str, object]) -> None:
    master = data["master_usd"]
    ffpi = data["ffpi_annual"]
    exposure = build_vulnerability(data["ghi"], data["worldbank"], float(controls["hunger_weight"]))

    st.title("Food Price Shock Early-Warning Portfolio")
    st.caption("Built for a UN Food Systems Summit Panel or regional food-security policy board.")
    story_insight(
        "Decision problem",
        "Producer-price and global food-price shocks become policy priorities when they meet countries with high hunger severity and high food-import dependency.",
        "Prioritise monitoring, preparedness funding, and import-buffer planning where exposure is highest.",
    )

    year_min = int(master["year"].min())
    year_max = int(master["year"].max())
    latest_ffpi = ffpi.dropna(subset=["ffpi_food"]).sort_values("year").iloc[-1]
    top_country = exposure.iloc[0]

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Producer-price coverage", f"{year_min}-{year_max}")
    c2.metric("AUS/NZ commodities", f"{master['item'].nunique():,}")
    c3.metric(f"Latest FFPI food ({int(latest_ffpi['year'])})", f"{latest_ffpi['ffpi_food']:.1f}")
    c4.metric("Highest exposure", top_country["country_ghi"], f"{top_country['exposure_index']:.1f}/100")

    st.caption(
        f"Current scope: {controls['year_range'][0]}-{controls['year_range'][1]}, "
        f"{controls['selected_item']}, {', '.join(controls['selected_countries'])}."
    )

    left, right = st.columns([0.54, 0.46])
    with left:
        st.subheader("Policy Watchlist")
        st.dataframe(priority_table(exposure, 10), width="stretch", hide_index=True)
        st.caption(
            "Suggested policy use is a triage label for discussion; it does not replace country-level field assessment."
        )

    with right:
        st.subheader("Narrative Route")
        st.markdown(
            """
            1. **Shock Context:** confirm global food-price stress periods.
            2. **Producer Signal:** inspect AUS/NZ commodity-price movement.
            3. **Vulnerability:** identify hunger-import exposure.
            4. **What-If Action:** stress-test scenario assumptions.
            5. **Evidence Base:** inspect cleaned data and volatility.
            """
        )
        st.warning(
            "Key boundary: this is an early-warning and prioritisation tool, not a bilateral trade-flow or causal import-price forecast."
        )


def slide_context(data: dict[str, pd.DataFrame], controls: dict[str, object]) -> None:
    ffpi = data["ffpi_annual"]
    ffpi_m = data["ffpi_monthly"]

    st.title("Why Food Price Shocks Matter")
    st.caption("Global food-price pressure has surged repeatedly; the question is who has the least capacity to absorb it.")
    story_insight(
        "Food price shocks are global, but vulnerability is not evenly shared",
        "The FFPI shows repeated global stress events across the last three decades.",
        "Use shock periods as context for interpreting producer-price signals and exposure priorities.",
    )

    selected_indices = list(controls["selected_indices"])
    year_range = controls["year_range"]

    ffpi_filtered = ffpi[ffpi["year"].between(year_range[0], year_range[1])]
    ffpi_m_filtered = ffpi_m[ffpi_m["year"].between(year_range[0], year_range[1])]

    st.subheader("Global Food Price Shock Timeline")
    fig = px.line(
        ffpi_filtered,
        x="year",
        y=selected_indices,
        color_discrete_map=FFPI_COLORS,
        labels={"value": "Index (2014-2016 = 100)", "variable": "Index"},
        title="When the World's Food Got Expensive",
        template=CHART_TEMPLATE,
    )
    fig.add_hline(y=100, line_dash="dash", line_color=COLORS["neutral"])
    add_crisis_bands(fig)
    apply_chart_style(fig, height=560)
    st.plotly_chart(fig, use_container_width=True)

    st.subheader("Supporting Evidence")
    c1, c2, c3 = st.columns(3)
    crisis_years = [2008, 2011, 2022]
    crisis = ffpi_filtered[ffpi_filtered["year"].isin(crisis_years)].set_index("year")
    for col, year in zip([c1, c2, c3], crisis_years):
        if year in crisis.index:
            col.metric(f"{year} FFPI food", f"{crisis.loc[year, 'ffpi_food']:.1f}")
        else:
            col.metric(f"{year} FFPI food", "outside range")

    fig_monthly = px.area(
        ffpi_m_filtered,
        x="date",
        y="ffpi_food",
        labels={"date": "Date", "ffpi_food": "FFPI food"},
        title="Monthly Food Price Shock Timeline",
        template=CHART_TEMPLATE,
    )
    fig_monthly.add_hline(y=100, line_dash="dash", line_color=COLORS["neutral"])
    fig_monthly.update_traces(line_color=COLORS["FFPI"], fillcolor="rgba(222, 73, 104, 0.16)")
    apply_chart_style(fig_monthly, height=390)
    st.plotly_chart(fig_monthly, use_container_width=True)


def slide_producer_signal(data: dict[str, pd.DataFrame], controls: dict[str, object]) -> None:
    master = data["master_usd"]
    master_idx = data["master_idx"]
    ffpi = data["ffpi_annual"]
    selected_item = str(controls["selected_item"])
    selected_countries = list(controls["selected_countries"])
    year_range = controls["year_range"]
    master_filtered = master[
        master["year"].between(year_range[0], year_range[1]) & master["country"].isin(selected_countries)
    ]
    ffpi_filtered = ffpi[ffpi["year"].between(year_range[0], year_range[1])]
    item_signal = master_filtered[master_filtered["item"].eq(selected_item)].merge(
        ffpi_filtered[["year", "ffpi_food"]], on="year", how="left", suffixes=("", "_annual")
    )
    if "ffpi_food_annual" in item_signal.columns:
        item_signal["ffpi_food"] = item_signal["ffpi_food_annual"].fillna(item_signal["ffpi_food"])
    item_signal = item_signal.sort_values(["country", "year"])

    st.title("Australia and New Zealand Producer Price Signal")
    st.caption("Commodity, country, and time range respond to the global analysis scope.")
    story_insight(
        "Producer prices provide the supply-side signal",
        f"{selected_item} is shown against the global FFPI to test whether the selected commodity moves through shock periods.",
        "Compare producer-price movement with global food-price pressure.",
    )
    if selected_item == "Wheat":
        st.caption(
            "Wheat is the default signal because it is a globally traded staple, appears in the AUS/NZ producer-price data, "
            "and provides an intuitive bridge between producer prices and food-security risk."
        )

    st.subheader(f"{selected_item} Producer Prices vs Global FFPI")
    if item_signal.empty:
        st.warning("The selected commodity/country/time combination is not available in the current producer-price file.")
    else:
        fig = make_subplots(specs=[[{"secondary_y": True}]])
        for country, color in [("Australia", COLORS["AUS"]), ("New Zealand", COLORS["NZL"])]:
            country_df = item_signal[item_signal["country"] == country]
            if country_df.empty:
                continue
            fig.add_trace(
                go.Scatter(
                    x=country_df["year"],
                    y=country_df["value"],
                    mode="lines+markers",
                    name=f"{country} {selected_item}",
                    line={"color": color, "width": 3},
                    hovertemplate="<b>%{fullData.name}</b><br>Year %{x}<br>USD/tonne: $%{y:,.0f}<extra></extra>",
                ),
                secondary_y=False,
            )
        ffpi_plot = ffpi_filtered[["year", "ffpi_food"]].dropna()
        fig.add_trace(
            go.Scatter(
                x=ffpi_plot["year"],
                y=ffpi_plot["ffpi_food"],
                mode="lines",
                name="Global FFPI",
                line={"color": COLORS["FFPI"], "width": 2, "dash": "dot"},
            ),
            secondary_y=True,
        )
        add_crisis_bands(fig)
        apply_chart_style(fig, height=560)
        fig.update_layout(title=f"AUS/NZ {selected_item} Producer Prices vs Global Food Price Index")
        fig.update_yaxes(title_text=f"{selected_item} producer price (USD/tonne)", secondary_y=False)
        fig.update_yaxes(title_text="FFPI food index", secondary_y=True)
        st.plotly_chart(fig, use_container_width=True)

    st.subheader("Supporting Evidence")
    corr = build_ffpi_correlation(
        master_idx[master_idx["year"].between(year_range[0], year_range[1])],
        ffpi_filtered,
    )
    fig_corr = px.bar(
        corr.sort_values(["country", "correlation"]),
        x="correlation",
        y="ffpi_index",
        color="country",
        barmode="group",
        orientation="h",
        color_discrete_map=COUNTRY_COLORS,
        labels={"correlation": "Pearson correlation", "ffpi_index": "FFPI sub-index"},
        title="AUS/NZ Producer Price Index Alignment with FFPI",
        template=CHART_TEMPLATE,
    )
    fig_corr.add_vline(x=0, line_color=COLORS["neutral"])
    apply_chart_style(fig_corr, height=390)
    st.plotly_chart(fig_corr, use_container_width=True)

def slide_vulnerability(data: dict[str, pd.DataFrame], controls: dict[str, object]) -> None:
    ghi = data["ghi"]
    worldbank = data["worldbank"]

    st.title("Who Is Most Vulnerable?")
    st.caption("Prioritise countries where food-access stress and food-import dependency reinforce each other.")
    story_insight(
        "Hunger severity needs an exposure channel",
        "Countries become priority cases when food-access stress and import dependency combine.",
        "The weighting can be adjusted for sensitivity testing while keeping the two weights constrained to 100%.",
    )

    hunger_weight = float(controls["hunger_weight"])
    import_weight = 1 - hunger_weight
    st.markdown(
        f"""
        **Exposure index**

        `({hunger_weight:.2f} x minmax(food-access stress) + {import_weight:.2f} x minmax(food import dependency)) x 100`

        **Vulnerability score**

        `{hunger_weight:.2f} x z(food-access stress) + {import_weight:.2f} x z(food import dependency)`
        """
    )

    exposure = build_corrected_vulnerability(data["ghi"], data["worldbank"], hunger_weight)

    st.subheader("Vulnerability Matrix")
    x_label = exposure["hunger_metric_label"].iloc[0] if "hunger_metric_label" in exposure else "GHI score 2025"
    fig = px.scatter(
        exposure,
        x="hunger_metric" if "hunger_metric" in exposure else "ghi_2025",
        y="food_import_pct",
        color="vulnerability_score",
        hover_name="country_ghi",
        hover_data={
            "food_import_pct": ":.1f",
            "vulnerability_score": ":.2f",
            "ghi_2025": ":.1f",
        },
        color_continuous_scale=RISK_SCALE,
        range_color=[
            exposure["vulnerability_score"].min(),
            exposure["vulnerability_score"].max(),
        ],
        labels={
            "hunger_metric": x_label,
            "food_import_pct": "Food imports (% of merchandise imports)",
            "vulnerability_score": "Vulnerability score",
        },
        title="Hunger/Food-Access Stress x Food Import Dependency",
        template=CHART_TEMPLATE,
    )
    fig.update_traces(marker={"size": 12, "opacity": 0.82, "line": {"color": "rgba(255,255,255,0.45)", "width": 0.7}})
    for _, row in exposure.nlargest(8, "vulnerability_score").iterrows():
        fig.add_annotation(
            x=row["hunger_metric"] if "hunger_metric" in row else row["ghi_2025"],
            y=row["food_import_pct"],
            text=row["country_ghi"],
            showarrow=False,
            xshift=8,
            yshift=4,
            font={"size": 10},
        )
    fig.add_vline(x=exposure["hunger_metric"].median(), line_dash="dash", line_color=COLORS["neutral"])
    fig.add_hline(y=exposure["food_import_pct"].median(), line_dash="dash", line_color=COLORS["neutral"])
    apply_chart_style(fig, height=610)
    st.plotly_chart(fig, use_container_width=True)

    st.subheader("Policy Watchlist")
    st.dataframe(priority_table(exposure, 10), width="stretch", hide_index=True)

    st.subheader("Sensitivity: Countries That Stay High Priority")
    sensitivity = build_weight_sensitivity(data["ghi"], data["worldbank"])
    if not sensitivity.empty:
        c1, c2 = st.columns([0.46, 0.54])
        with c1:
            st.dataframe(
                sensitivity.head(10).round({"average_rank": 1, "average_exposure": 1}).rename(
                    columns={
                        "country": "Country",
                        "appearances": "Top-10 appearances",
                        "best_rank": "Best rank",
                        "average_rank": "Average rank",
                        "average_exposure": "Average exposure",
                    }
                ),
                width="stretch",
                hide_index=True,
            )
        with c2:
            fig_stability = px.bar(
                sensitivity.head(10).sort_values("appearances"),
                x="appearances",
                y="country",
                orientation="h",
                color="average_exposure",
                color_continuous_scale=RISK_SCALE,
                labels={
                    "appearances": "Top-10 appearances across tested weights",
                    "country": "",
                    "average_exposure": "Average exposure",
                },
                title="Priority Robustness across Food-Access Weights 0.40-0.80",
                template=CHART_TEMPLATE,
            )
            apply_chart_style(fig_stability, height=390)
            st.plotly_chart(fig_stability, use_container_width=True)

    st.subheader("Supporting Evidence")
    col1, col2 = st.columns(2)
    with col1:
        ghi_valid = ghi.dropna(subset=["iso3", "ghi_2025"])
        fig_map = px.choropleth(
            ghi_valid,
            locations="iso3",
            color="ghi_2025",
            hover_name="country_ghi",
            color_continuous_scale=RISK_SCALE,
            labels={"ghi_2025": "GHI 2025"},
            title="Global Hunger Index 2025",
            template=CHART_TEMPLATE,
        )
        apply_chart_style(fig_map, height=420)
        fig_map.update_layout(geo_showframe=False, geo_showcoastlines=True, geo_coastlinecolor=COLORS["neutral"])
        st.plotly_chart(fig_map, use_container_width=True)
    with col2:
        latest_imports = (
            worldbank.dropna(subset=["food_import_pct"])
            .sort_values("year", ascending=False)
            .groupby("iso3", as_index=False)
            .first()
        )
        fig_hist = px.histogram(
            latest_imports,
            x="food_import_pct",
            nbins=35,
            color_discrete_sequence=[COLORS["AUS"]],
            labels={"food_import_pct": "Food imports (% of merchandise imports)"},
            title="Distribution of Food Import Dependency",
            template=CHART_TEMPLATE,
        )
        fig_hist.add_vline(x=latest_imports["food_import_pct"].median(), line_dash="dash", line_color=COLORS["neutral"])
        apply_chart_style(fig_hist, height=420, show_legend=False)
        st.plotly_chart(fig_hist, use_container_width=True)


def slide_what_if(data: dict[str, pd.DataFrame], controls: dict[str, object]) -> None:
    exposure = build_corrected_vulnerability(data["ghi"], data["worldbank"], float(controls["hunger_weight"]))

    st.title("What-If Scenario and Priority Ranking")
    st.caption(
        "Stress-test how an AUS/NZ producer-price shock could translate into priority monitoring for structurally exposed countries."
    )
    story_insight(
        "The scenario converts insight into action",
        "The ranking shows where early-warning monitoring should focus under an illustrative AUS/NZ producer-price shock.",
        "Assumptions are adjustable so the ranking can be stress-tested, not treated as a single-point forecast.",
    )

    shock_pct = float(controls["shock_pct"])
    basket_share = float(controls["basket_share"])
    transmission_coeff = float(controls["transmission_coeff"])
    pass_through_rate = float(controls["pass_through_rate"])
    scenario, effective_global_pressure_pct = build_scenario(
        exposure,
        shock_pct,
        pass_through_rate,
        basket_share,
        transmission_coeff,
    )

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Producer shock", f"+{shock_pct}%")
    c2.metric("Basket share", f"{basket_share:.0%}")
    c3.metric("Transmission x pass-through", f"{transmission_coeff:.0%} x {pass_through_rate:.0%}")
    c4.metric("Effective global pressure", f"{effective_global_pressure_pct:.2f}%")

    st.subheader("Scenario Priority Ranking")
    top = scenario.nlargest(12, "scenario_pressure_score")
    max_scenario, _ = build_scenario(exposure, 60, 1.0, 1.0, 1.0)
    max_scenario_score = max(max_scenario["scenario_pressure_score"].max(), 0.01)
    fig = px.bar(
        top,
        x="scenario_pressure_score",
        y="country_ghi",
        orientation="h",
        color="scenario_pressure_score",
        color_continuous_scale=RISK_SCALE,
        range_color=[0, max_scenario_score],
        hover_data={
            "hunger_metric": ":.1f",
            "food_import_pct": ":.1f",
            "vulnerability_score": ":.2f",
            "implied_import_cost_pressure_pct": ":.2f",
            "scenario_pressure_score": ":.2f",
        },
        labels={"country_ghi": "", "scenario_pressure_score": "Scenario pressure score"},
        title=f"Priority Countries under +{shock_pct}% Producer Shock",
        template=CHART_TEMPLATE,
    )
    apply_chart_style(fig, height=560)
    fig.update_layout(coloraxis_showscale=False)
    fig.update_xaxes(range=[0, max_scenario_score * 1.08])
    fig.update_yaxes(categoryorder="total ascending")
    st.plotly_chart(fig, use_container_width=True)

    st.subheader("Recommended Policy Use")
    st.dataframe(priority_table(scenario, 10), width="stretch", hide_index=True)

    with st.expander("Scenario assumptions and boundaries", expanded=False):
        st.markdown(
            f"""
            **Formula**

            `effective_global_pressure = producer_shock x basket_share x transmission_coeff x pass_through_rate`

            **Current assumptions**

            - Relevant import basket share: `{basket_share:.2f}`
            - Transmission coefficient: `{transmission_coeff:.2f}`
            - Pass-through rate: `{pass_through_rate:.2f}` from the slider

            This remains an illustrative prioritisation model. It does not include bilateral trade flows,
            exchange rates, freight costs, tariffs, subsidies, trade margins, or supplier substitution.
            """
        )


def appendix(data: dict[str, pd.DataFrame], controls: dict[str, object]) -> None:
    st.title("Evidence Base")
    st.caption("Transparent source tables and volatility checks supporting the policy narrative.")

    st.subheader("Dataset Explorer")
    label_to_key = {
        "Producer prices, USD": "master_usd",
        "Producer price index": "master_idx",
        "FFPI annual": "ffpi_annual",
        "FFPI monthly": "ffpi_monthly",
        "Global Hunger Index": "ghi",
        "World Bank food imports": "worldbank",
    }
    label = str(controls["dataset_label"])
    df = data[label_to_key[label]]
    st.caption(f"{len(df):,} rows x {df.shape[1]:,} columns")
    st.dataframe(df, width="stretch", hide_index=True)

    st.subheader("Commodity Volatility")
    country = str(controls["volatility_country"])
    subset = data["master_usd"][data["master_usd"]["country"] == country]
    vol = subset.groupby("item")["value"].agg(["mean", "std", "count"]).reset_index()
    vol = vol[vol["count"] >= 5]
    vol["cv_pct"] = vol["std"] / vol["mean"] * 100
    vol = vol.sort_values("cv_pct", ascending=False).head(15)
    fig_vol = px.bar(
        vol.sort_values("cv_pct"),
        x="cv_pct",
        y="item",
        orientation="h",
        color="cv_pct",
        color_continuous_scale=RISK_SCALE,
        labels={"cv_pct": "Coefficient of variation (%)", "item": ""},
        title=f"{country}: Most Volatile Producer Prices",
        template=CHART_TEMPLATE,
    )
    fig_vol.add_vline(x=50, line_dash="dash", line_color=COLORS["risk"], annotation_text="50% CV")
    apply_chart_style(fig_vol, height=520)
    fig_vol.update_layout(coloraxis_showscale=False)
    st.plotly_chart(fig_vol, use_container_width=True)


def archived_pitch_brief(data: dict[str, pd.DataFrame]) -> None:
    st.title("Archived Narrative Brief")
    st.caption("A presentation-ready path through the dashboard: what changed, who is exposed, and what to do next.")

    ffpi = data["ffpi_annual"]
    wheat = build_wheat_signal(data["master_usd"], ffpi)
    exposure = build_undernourishment_exposure(data["ghi"], data["worldbank"])

    st.subheader("1. The Shock Context: Three Global Price Surges")
    st.write(
        "This visual supports the opening claim: food price shocks are global, but their consequences are uneven."
    )
    fig_ffpi = px.line(
        ffpi[ffpi["year"] >= 1991],
        x="year",
        y=["ffpi_food", "ffpi_cereals", "ffpi_oils", "ffpi_sugar"],
        color_discrete_map=FFPI_COLORS,
        labels={"value": "Index (2014-2016 = 100)", "variable": "Index"},
        title="FAO Food Price Index: shock periods since 1991",
        template=CHART_TEMPLATE,
    )
    fig_ffpi.add_hline(y=100, line_dash="dash", line_color=COLORS["neutral"])
    add_crisis_bands(fig_ffpi)
    apply_chart_style(fig_ffpi, height=430)
    st.plotly_chart(fig_ffpi, use_container_width=True)

    st.subheader("2. The Local Signal: Wheat Connects AUS/NZ to the Global Story")
    st.write(
        "We added this as a pitch visual because wheat is easy to understand, globally important, and visible in the crisis periods."
    )
    if not wheat.empty:
        fig_wheat = make_subplots(specs=[[{"secondary_y": True}]])
        for country, color in [("Australia", COLORS["AUS"]), ("New Zealand", COLORS["NZL"])]:
            country_df = wheat[wheat["country"] == country]
            if country_df.empty:
                continue
            fig_wheat.add_trace(
                go.Scatter(
                    x=country_df["year"],
                    y=country_df["value"],
                    mode="lines+markers",
                    name=f"{country} wheat",
                    line={"color": color, "width": 3},
                    hovertemplate="<b>%{fullData.name}</b><br>Year %{x}<br>USD/tonne: $%{y:,.0f}<extra></extra>",
                ),
                secondary_y=False,
            )
        ffpi_plot = ffpi[["year", "ffpi_food"]].dropna()
        fig_wheat.add_trace(
            go.Scatter(
                x=ffpi_plot["year"],
                y=ffpi_plot["ffpi_food"],
                mode="lines",
                name="Global FFPI",
                line={"color": COLORS["FFPI"], "width": 2, "dash": "dot"},
            ),
            secondary_y=True,
        )
        add_crisis_bands(fig_wheat)
        apply_chart_style(fig_wheat, height=470)
        fig_wheat.update_layout(title="Wheat Producer Prices vs Global Food Price Index")
        fig_wheat.update_yaxes(title_text="Wheat producer price (USD/tonne)", secondary_y=False)
        fig_wheat.update_yaxes(title_text="FFPI food index", secondary_y=True)
        st.plotly_chart(fig_wheat, use_container_width=True)
    else:
        st.info("Wheat is not available in the current producer-price file.")

    st.subheader("3. The Human Lens: Hunger-Import Exposure")
    st.write(
        "This improves the deck's vulnerability argument by using undernourishment where available, not just the broad GHI score."
    )
    x_label = exposure["hunger_metric_label"].iloc[0] if "hunger_metric_label" in exposure else "GHI score 2025"
    fig_exposure = px.scatter(
        exposure,
        x="hunger_metric" if "hunger_metric" in exposure else "ghi_2025",
        y="food_import_pct",
        color="exposure_index",
        size="exposure_index",
        size_max=24,
        hover_name="country_ghi",
        hover_data={"food_import_pct": ":.1f", "exposure_index": ":.1f", "ghi_2025": ":.1f"},
        color_continuous_scale=RISK_SCALE,
        labels={
            "hunger_metric": x_label,
            "food_import_pct": "Food imports (% of merchandise imports)",
            "exposure_index": "Exposure index",
        },
        title="Countries Facing Both Food-Access Stress and Import Exposure",
        template=CHART_TEMPLATE,
    )
    top_labels = exposure.nlargest(8, "exposure_index")
    for _, row in top_labels.iterrows():
        fig_exposure.add_annotation(
            x=row["hunger_metric"] if "hunger_metric" in row else row["ghi_2025"],
            y=row["food_import_pct"],
            text=row["country_ghi"],
            showarrow=False,
            xshift=8,
            yshift=4,
            font={"size": 10},
        )
    apply_chart_style(fig_exposure, height=560)
    st.plotly_chart(fig_exposure, use_container_width=True)

    st.subheader("4. The Action: Scenario Priority List")
    shock_pct = st.slider("Presentation shock scenario", min_value=0, max_value=60, value=20, step=5, format="+%d%%")
    pass_through_pct = st.slider("Presentation pass-through assumption", 0, 100, 25, step=5, format="%d%%")
    scenario = exposure.copy()
    global_pressure = shock_pct * pass_through_pct / 100
    scenario["implied_import_cost_pressure_pct"] = scenario["food_import_pct"] * global_pressure / 100
    scenario["scenario_pressure_score"] = scenario["exposure_index"] * scenario["implied_import_cost_pressure_pct"] / 100
    top = scenario.nlargest(12, "scenario_pressure_score").sort_values("scenario_pressure_score")
    fig_priority = px.bar(
        top,
        x="scenario_pressure_score",
        y="country_ghi",
        orientation="h",
        color="scenario_pressure_score",
        color_continuous_scale=RISK_SCALE,
        labels={"country_ghi": "", "scenario_pressure_score": "Scenario pressure score"},
        title=f"Priority Countries under +{shock_pct}% Producer Shock x {pass_through_pct}% Pass-through",
        template=CHART_TEMPLATE,
    )
    apply_chart_style(fig_priority, height=470)
    fig_priority.update_layout(coloraxis_showscale=False)
    st.plotly_chart(fig_priority, use_container_width=True)

    with st.expander("Why these visuals support the narrative"):
        st.markdown(
            """
            - **FFPI shock context** gives the audience the global problem in one screen.
            - **Wheat signal** turns AUS/NZ producer prices into a memorable indicator commodity.
            - **Undernourishment exposure** makes the human-centred design stronger than a generic risk score.
            - **Scenario priority list** converts the analysis into an action: who should be monitored first.
            """
        )


def overview(data: dict[str, pd.DataFrame]) -> None:
    master = data["master_usd"]
    ffpi = data["ffpi_annual"]
    vuln = build_vulnerability(data["ghi"], data["worldbank"])

    st.title("FAOSTAT Producer Prices: Shock Exposure Dashboard")
    st.caption("Australia and New Zealand producer prices, global food price shocks, and hunger-import exposure.")
    st.write(
        "Follow the story from supply-side producer prices, to global food price context, "
        "to countries least able to absorb price pressure."
    )
    story_insight(
        "Start with the map of evidence",
        "The dashboard joins producer prices, FFPI, hunger, and food-import dependency so the audience can move from prices to people.",
        "Use this page to establish scope before diving into the causal caveats.",
    )

    year_min = int(master["year"].min())
    year_max = int(master["year"].max())
    latest_ffpi = ffpi.dropna(subset=["ffpi_food"]).sort_values("year").iloc[-1]

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Coverage", f"{year_min}-{year_max}")
    c2.metric("Commodities", f"{master['item'].nunique():,}")
    c3.metric("Producer price rows", f"{len(master):,}")
    c4.metric(f"Latest FFPI food ({int(latest_ffpi['year'])})", f"{latest_ffpi['ffpi_food']:.1f}")

    country_summary = (
        master.groupby(["iso3", "country"], as_index=False)
        .agg(avg_usd_per_tonne=("value", "mean"), commodities=("item", "nunique"), outliers=("is_outlier", "sum"))
        .sort_values("iso3")
    )
    country_summary["outlier_pct"] = country_summary["outliers"] / country_summary["commodities"].clip(lower=1) * 100
    country_summary["avg_usd_per_tonne"] = country_summary["avg_usd_per_tonne"].round(0)

    left, right = st.columns([0.42, 0.58])
    with left:
        st.subheader("Country Coverage Snapshot")
        st.dataframe(
            country_summary[
                ["country", "commodities", "avg_usd_per_tonne", "outliers"]
            ].rename(
                columns={
                    "country": "Country",
                    "commodities": "Commodities",
                    "avg_usd_per_tonne": "Avg USD/tonne",
                    "outliers": "Flagged outliers",
                }
            ),
            width="stretch",
            hide_index=True,
        )
        st.caption("Compact overview only. The detailed price story starts in `2. Producer Signal` and `5. Evidence Base`.")

    with right:
        st.subheader("Highest Hunger-Import Exposure")
        st.dataframe(
            vuln.head(8)[["country_ghi", "ghi_2025", "food_import_pct", "exposure_index"]]
            .round(2)
            .rename(
                columns={
                    "country_ghi": "Country",
                    "ghi_2025": "GHI 2025",
                    "food_import_pct": "Food imports %",
                    "exposure_index": "Exposure index",
                }
            ),
            width="stretch",
            hide_index=True,
        )
    st.caption(
        "This summary establishes dataset scope and priority exposure signals before the detailed analytical views."
    )


def price_trends(data: dict[str, pd.DataFrame]) -> None:
    master = data["master_usd"]
    ffpi = data["ffpi_annual"]

    st.title("Producer Price Trends")
    st.caption("Start with the supply-side signal: how AUS/NZ commodity prices move through time.")
    story_insight(
        "What is changing?",
        "Producer prices show the supply-side signal, while the FFPI overlay shows whether the selected commodity moves with global food-price pressure.",
        "Start with Wheat for the clearest presentation example, then test other commodities interactively.",
    )
    items = sorted(master["item"].dropna().unique())
    default_item = "Wheat" if "Wheat" in items else items[0]
    selected_item = st.selectbox("Commodity", items, index=items.index(default_item))

    countries = st.multiselect(
        "Country",
        options=["Australia", "New Zealand"],
        default=["Australia", "New Zealand"],
    )

    filtered = master[(master["item"] == selected_item) & (master["country"].isin(countries))].sort_values("year")
    fig = make_subplots(specs=[[{"secondary_y": True}]])

    for country, color in [("Australia", COLORS["AUS"]), ("New Zealand", COLORS["NZL"])]:
        country_df = filtered[filtered["country"] == country]
        if country_df.empty:
            continue
        fig.add_trace(
            go.Scatter(
                x=country_df["year"],
                y=country_df["value"],
                mode="lines+markers",
                name=country,
                line={"color": color, "width": 3},
                customdata=country_df[["flag", "is_imputed", "is_outlier"]],
                hovertemplate=(
                    "<b>%{fullData.name}</b><br>Year %{x}<br>"
                    "USD/tonne: $%{y:,.0f}<br>Flag: %{customdata[0]}<br>"
                    "Imputed: %{customdata[1]}<br>Outlier: %{customdata[2]}<extra></extra>"
                ),
            ),
            secondary_y=False,
        )

    ffpi_plot = ffpi[["year", "ffpi_food"]].dropna()
    fig.add_trace(
        go.Scatter(
            x=ffpi_plot["year"],
            y=ffpi_plot["ffpi_food"],
            mode="lines",
            name="FFPI food",
            line={"color": COLORS["FFPI"], "width": 2, "dash": "dot"},
        ),
        secondary_y=True,
    )
    add_crisis_bands(fig)
    apply_chart_style(fig, height=580)
    fig.update_layout(title=f"{selected_item}: AUS/NZ Producer Price vs Global FFPI")
    fig.update_xaxes(title_text="Year")
    fig.update_yaxes(title_text="Producer price (USD/tonne)", secondary_y=False)
    fig.update_yaxes(title_text="FFPI food index", secondary_y=True)
    st.plotly_chart(fig, use_container_width=True)

    outliers = filtered[filtered["is_outlier"].astype(str).str.lower() == "true"]
    if not outliers.empty:
        st.subheader("Flagged Outlier Observations")
        st.dataframe(
            outliers[["country", "item", "year", "value", "flag_description"]].sort_values("year"),
            width="stretch",
            hide_index=True,
        )


def volatility(data: dict[str, pd.DataFrame]) -> None:
    master = data["master_usd"]
    st.title("Commodity Volatility")
    st.caption("Volatility highlights which commodities have unstable producer prices and may deserve closer monitoring.")
    story_insight(
        "Which commodities deserve attention?",
        "Average price alone is not enough; coefficient of variation highlights commodities with unstable producer prices after normalising for scale.",
        "Use high-volatility commodities as candidates for early-warning monitoring.",
    )

    country = st.radio("Country", ["Australia", "New Zealand"], horizontal=True)
    top_n = st.slider("Number of commodities", min_value=5, max_value=30, value=15, step=5)
    subset = master[master["country"] == country]

    vol = subset.groupby("item")["value"].agg(["mean", "std", "count"]).reset_index()
    vol = vol[vol["count"] >= 5]
    vol["cv_pct"] = vol["std"] / vol["mean"] * 100
    vol = vol.sort_values("cv_pct", ascending=False).head(top_n)

    fig = px.bar(
        vol.sort_values("cv_pct"),
        x="cv_pct",
        y="item",
        orientation="h",
        color="cv_pct",
        color_continuous_scale=RISK_SCALE,
        labels={"cv_pct": "Coefficient of variation (%)", "item": ""},
        hover_data={"mean": ":,.0f", "std": ":,.0f", "count": True},
        title=f"{country}: Most Volatile Producer Prices",
        template=CHART_TEMPLATE,
    )
    fig.add_vline(x=50, line_dash="dash", line_color=COLORS["risk"], annotation_text="50% CV")
    apply_chart_style(fig, height=max(500, top_n * 28))
    fig.update_layout(coloraxis_showscale=False)
    st.plotly_chart(fig, use_container_width=True)


def global_context(data: dict[str, pd.DataFrame]) -> None:
    ffpi_a = data["ffpi_annual"]
    ffpi_m = data["ffpi_monthly"]
    ghi = data["ghi"]
    master_idx = data["master_idx"]

    st.title("Global Shock and Hunger Context")
    st.caption("Connect local producer-price movement to global food price shock periods and hunger severity.")
    story_insight(
        "Global context",
        "FFPI sub-indices reveal the shock environment, while GHI shows why equal price pressure can create unequal human consequences.",
        "Connect the producer-price signal to global food-security risk.",
    )
    ffpi_cols = ["ffpi_food", "ffpi_cereals", "ffpi_meat", "ffpi_dairy", "ffpi_oils", "ffpi_sugar"]
    selected = st.multiselect(
        "FFPI indices",
        ffpi_cols,
        default=["ffpi_food", "ffpi_cereals", "ffpi_oils"],
        format_func=lambda x: x.replace("ffpi_", "").title(),
    )
    if not selected:
        selected = ["ffpi_food"]

    fig = px.line(
        ffpi_a[ffpi_a["year"] >= 1991],
        x="year",
        y=selected,
        color_discrete_map=FFPI_COLORS,
        labels={"value": "Index (2014-2016 = 100)", "year": "Year", "variable": "Index"},
        title="FAO Food Price Index Sub-Indices",
        template=CHART_TEMPLATE,
    )
    fig.add_hline(y=100, line_dash="dash", line_color=COLORS["neutral"])
    add_crisis_bands(fig)
    apply_chart_style(fig, height=520)
    st.plotly_chart(fig, use_container_width=True)

    st.subheader("AUS/NZ Price Index Alignment with FFPI")
    corr = build_ffpi_correlation(master_idx, ffpi_a)
    fig_corr = px.bar(
        corr.sort_values(["country", "correlation"]),
        x="correlation",
        y="ffpi_index",
        color="country",
        barmode="group",
        orientation="h",
        color_discrete_map=COUNTRY_COLORS,
        labels={"correlation": "Pearson correlation", "ffpi_index": "FFPI sub-index"},
        title="Correlation between AUS/NZ Producer Price Index and FFPI Sub-Indices",
        template=CHART_TEMPLATE,
    )
    fig_corr.add_vline(x=0, line_color=COLORS["neutral"])
    apply_chart_style(fig_corr, height=420)
    st.plotly_chart(fig_corr, use_container_width=True)

    crisis_years = [2008, 2011, 2022]
    crisis = ffpi_a[ffpi_a["year"].isin(crisis_years)].set_index("year")
    st.subheader("Crisis-Year Price Shock Markers")
    cols = st.columns(len(crisis_years))
    for col, year in zip(cols, crisis_years):
        if year in crisis.index:
            col.metric(f"{year} FFPI food", f"{crisis.loc[year, 'ffpi_food']:.1f}")

    monthly = ffpi_m[ffpi_m["year"] >= 2000].copy()
    fig_monthly = px.area(
        monthly,
        x="date",
        y="ffpi_food",
        labels={"date": "Date", "ffpi_food": "FFPI food"},
        title="Monthly Food Price Shock Timeline",
        template=CHART_TEMPLATE,
    )
    fig_monthly.add_hline(y=100, line_dash="dash", line_color=COLORS["neutral"])
    fig_monthly.update_traces(line_color=COLORS["FFPI"], fillcolor="rgba(222, 73, 104, 0.16)")
    apply_chart_style(fig_monthly, height=420)
    st.plotly_chart(fig_monthly, use_container_width=True)

    ghi_valid = ghi.dropna(subset=["iso3", "ghi_2025"])
    fig_map = px.choropleth(
        ghi_valid,
        locations="iso3",
        color="ghi_2025",
        hover_name="country_ghi",
        color_continuous_scale=RISK_SCALE,
        labels={"ghi_2025": "GHI 2025"},
        title="Global Hunger Index 2025",
        template=CHART_TEMPLATE,
    )
    apply_chart_style(fig_map, height=520)
    fig_map.update_layout(geo_showframe=False, geo_showcoastlines=True, geo_coastlinecolor=COLORS["neutral"])
    st.plotly_chart(fig_map, use_container_width=True)


def vulnerability(data: dict[str, pd.DataFrame]) -> None:
    vuln = build_vulnerability(data["ghi"], data["worldbank"])
    st.title("Hunger-Import Exposure Matrix")
    st.caption("A policy-board view of countries where food-access stress and import dependency reinforce each other.")
    story_insight(
        "Who is least able to absorb the shock?",
        "Countries become priority cases when hunger severity and food-import dependency combine; the sliders then translate that baseline exposure into a scenario ranking.",
        "This is a prioritisation tool, not a bilateral trade-flow forecast.",
    )

    shock_pct = st.slider("Producer price shock", min_value=0, max_value=60, value=20, step=5, format="+%d%%")
    pass_through_pct = st.slider(
        "Producer-to-import pass-through", min_value=0, max_value=100, value=50, step=5, format="%d%%"
    )
    basket_share = SCENARIO_ASSUMPTIONS["basket_share"]
    transmission_coeff = SCENARIO_ASSUMPTIONS["transmission_coeff"]
    pass_through_rate = pass_through_pct / 100
    vuln = vuln.copy()
    effective_global_pressure_pct = shock_pct * basket_share * transmission_coeff * pass_through_rate
    vuln["implied_import_cost_pressure_pct"] = vuln["food_import_pct"] * effective_global_pressure_pct / 100
    vuln["scenario_pressure_score"] = vuln["exposure_index"] * vuln["implied_import_cost_pressure_pct"] / 100

    highest = vuln.iloc[0]
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Highest exposure country", highest["country_ghi"])
    c2.metric("Exposure index", f"{highest['exposure_index']:.1f}/100")
    c3.metric("Producer shock", f"+{shock_pct}%")
    c4.metric("Pass-through", f"{pass_through_rate:.0%}")
    c5.metric("Effective global pressure", f"{effective_global_pressure_pct:.2f}%")

    st.caption(
        "This is an exposure index, not a bilateral AUS/NZ trade-flow model. "
        "Dots stay in the same x/y position because GHI and food import dependency are baseline indicators; "
        "the shock slider changes the illustrative scenario pressure shown by bubble size, colour, and bar length. "
        "The scenario uses the baseline assumptions: basket share 45% and transmission coefficient 0.429; "
        "pass-through is adjustable because policy buffers, exchange rates, freight, margins, and substitution affect final import costs."
    )

    with st.expander("Methodology and caveats", expanded=False):
        st.markdown(
            """
            **Exposure index**

            `exposure_index = (0.60 x minmax(GHI 2025) + 0.40 x minmax(food import dependency)) x 100`

            **Scenario pressure**

            `effective_global_pressure = producer_price_shock x basket_share x transmission_coeff x pass_through_rate`

            `scenario_pressure_score = exposure_index x implied_import_cost_pressure / 100`

            **Current assumptions**

            - `basket_share = 0.45`
            - `transmission_coeff = 0.429`
            - `pass_through_rate = user-selected slider`

            This is not a bilateral AUS/NZ trade-flow model. It does not estimate exchange-rate effects,
            freight costs, tariffs, subsidies, or supplier substitution. It is a prioritisation tool for
            identifying countries structurally exposed to global food-price stress.
            """
        )

    x_med = vuln["ghi_2025"].median()
    y_med = vuln["food_import_pct"].median()
    max_possible_global_pressure_pct = 60 * basket_share * transmission_coeff * 1.0
    max_scenario_pressure = max(data["worldbank"]["food_import_pct"].max() * max_possible_global_pressure_pct / 100, 1)
    vuln["scenario_bubble_size"] = vuln["scenario_pressure_score"].clip(lower=0.02)
    max_scenario_score = max(
        (vuln["exposure_index"] * vuln["food_import_pct"] * max_possible_global_pressure_pct / 10000).max(),
        1,
    )
    fig = px.scatter(
        vuln,
        x="ghi_2025",
        y="food_import_pct",
        size="scenario_bubble_size",
        size_max=28,
        color="implied_import_cost_pressure_pct",
        hover_name="country_ghi",
        hover_data={
            "ghi_2025": ":.1f",
            "food_import_pct": ":.1f",
            "exposure_index": ":.1f",
            "scenario_bubble_size": False,
            "implied_import_cost_pressure_pct": ":.2f",
            "scenario_pressure_score": ":.2f",
        },
        color_continuous_scale=RISK_SCALE,
        range_color=[0, max_scenario_pressure],
        labels={
            "ghi_2025": "GHI score 2025",
            "food_import_pct": "Food imports (% of merchandise imports)",
            "exposure_index": "Hunger-import exposure index",
            "scenario_bubble_size": "Scenario pressure score",
            "implied_import_cost_pressure_pct": "Implied import-cost pressure (%)",
            "scenario_pressure_score": "Scenario pressure score",
        },
        title=(
            "Hunger Severity x Food Import Dependency "
            f"under +{shock_pct}% Producer Shock and {effective_global_pressure_pct:.2f}% Effective Global Pressure"
        ),
        template=CHART_TEMPLATE,
    )
    fig.add_vline(x=x_med, line_dash="dash", line_color=COLORS["neutral"])
    fig.add_hline(y=y_med, line_dash="dash", line_color=COLORS["neutral"])
    fig.update_traces(marker={"line": {"color": "rgba(255,255,255,0.42)", "width": 0.8}, "opacity": 0.9})
    apply_chart_style(fig, height=620)
    st.plotly_chart(fig, use_container_width=True)

    top = vuln.nlargest(15, "scenario_pressure_score").sort_values("scenario_pressure_score")
    fig_bar = px.bar(
        top,
        x="scenario_pressure_score",
        y="country_ghi",
        orientation="h",
        color="scenario_pressure_score",
        color_continuous_scale=RISK_SCALE,
        range_color=[0, max_scenario_score],
        hover_data={
            "ghi_2025": ":.1f",
            "food_import_pct": ":.1f",
            "exposure_index": ":.1f",
            "implied_import_cost_pressure_pct": ":.2f",
        },
        labels={
            "scenario_pressure_score": "Scenario pressure score",
            "country_ghi": "",
            "exposure_index": "Exposure index",
        },
        title=(
            "Illustrative What-if Ranking: "
            f"+{shock_pct}% Producer Shock x Basket/Transmission/Pass-through Assumptions"
        ),
        template=CHART_TEMPLATE,
    )
    fig_bar.update_traces(marker_line_color="rgba(255,255,255,0.22)", marker_line_width=0.8)
    fig_bar.update_xaxes(range=[0, max_scenario_score])
    apply_chart_style(fig_bar, height=560)
    st.plotly_chart(fig_bar, use_container_width=True)

    st.subheader("Policy Priority Actions")
    st.dataframe(priority_table(vuln, 10), width="stretch", hide_index=True)

    st.subheader("Country Deep Dive")
    country = st.selectbox("Inspect a country", vuln["country_ghi"].tolist())
    selected = vuln[vuln["country_ghi"] == country].iloc[0]
    d1, d2, d3, d4 = st.columns(4)
    d1.metric("GHI 2025", f"{selected['ghi_2025']:.1f}")
    d2.metric("Food import dependency", f"{selected['food_import_pct']:.1f}%")
    d3.metric("Exposure index", f"{selected['exposure_index']:.1f}/100")
    d4.metric("Scenario pressure", f"{selected['scenario_pressure_score']:.2f}")

    trend_cols = ["ghi_2000", "ghi_2008", "ghi_2016", "ghi_2025"]
    trend = pd.DataFrame(
        {
            "year": [2000, 2008, 2016, 2025],
            "ghi_score": [selected[col] for col in trend_cols],
        }
    ).dropna()
    fig_trend = px.line(
        trend,
        x="year",
        y="ghi_score",
        markers=True,
        labels={"year": "Year", "ghi_score": "GHI score"},
        title=f"{country}: Hunger Score Trend",
        template=CHART_TEMPLATE,
    )
    fig_trend.update_traces(line_color=COLORS["risk"], marker_color=COLORS["risk"])
    apply_chart_style(fig_trend, height=340)
    st.plotly_chart(fig_trend, use_container_width=True)


def data_explorer(data: dict[str, pd.DataFrame]) -> None:
    st.title("Data Explorer")
    story_insight(
        "Can the audience inspect the evidence?",
        "This tab exposes the cleaned source tables so the dashboard remains transparent and auditable.",
        "Use this view when a viewer wants to inspect where a number or variable came from.",
    )
    label_to_key = {
        "Producer prices, USD": "master_usd",
        "Producer price index": "master_idx",
        "FFPI annual": "ffpi_annual",
        "FFPI monthly": "ffpi_monthly",
        "Global Hunger Index": "ghi",
        "World Bank food imports": "worldbank",
    }
    label = st.selectbox("Dataset", list(label_to_key))
    df = data[label_to_key[label]]
    st.caption(f"{len(df):,} rows x {df.shape[1]:,} columns")
    st.dataframe(df, width="stretch", hide_index=True)


def main() -> None:
    st.set_page_config(page_title="Food Price Shock Early Warning", layout="wide")
    data = load_data()

    if not require_data(data):
        return

    controls = sidebar_controls(data)

    tabs = st.tabs(
        [
            "Executive Brief",
            "1. Shock Context",
            "2. Producer Signal",
            "3. Vulnerability",
            "4. What-If Action",
            "5. Evidence Base",
        ]
    )
    with tabs[0]:
        executive_brief(data, controls)
    with tabs[1]:
        slide_context(data, controls)
    with tabs[2]:
        slide_producer_signal(data, controls)
    with tabs[3]:
        slide_vulnerability(data, controls)
    with tabs[4]:
        slide_what_if(data, controls)
    with tabs[5]:
        appendix(data, controls)


if __name__ == "__main__":
    main()
